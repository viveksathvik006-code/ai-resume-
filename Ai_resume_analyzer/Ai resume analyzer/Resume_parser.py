import re

from pypdf import PdfReader
from docx import Document


def read_pdf(file):
    """Return the text from every page of a PDF."""
    reader = PdfReader(file)
    pages = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages.append(text)
    return "\n".join(pages)


def read_docx(file):
    """Return the text from every paragraph (and table) of a DOCX file."""
    document = Document(file)
    lines = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            lines.append(paragraph.text)

    # Some resumes keep their skills inside a table, so read those as well
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    lines.append(cell.text)

    return "\n".join(lines)


def extract_text(uploaded_file):
    """Pick the right reader based on the file name."""
    name = uploaded_file.name.lower()

    if name.endswith(".pdf"):
        return read_pdf(uploaded_file)
    if name.endswith(".docx"):
        return read_docx(uploaded_file)

    raise ValueError("Please upload a PDF or DOCX file.")


def clean_text(text):
    """
    Make the text easy to compare:
    lowercase it, drop emails and links, remove odd symbols and extra spaces.
    The symbols + # . / - are kept so that C++, C#, .NET, Node.js,
    CI/CD and scikit-learn still make sense.
    """
    text = text.lower()
    text = re.sub(r"\S+@\S+", " ", text)                 
    text = re.sub(r"(https?://|www\.)\S+", " ", text)    
    text = re.sub(r"[^a-z0-9+#./\-\s]", " ", text)       
    text = re.sub(r"\s+", " ", text)                     
    return text.strip()
